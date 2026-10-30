"""전국 키워드(31,023) 중 서울·경기 8,070개 선정 → data/site-pages.json + keywords/커튼블라인드_서울경기_8000.xlsx

  python -I scripts/select_seoul_gg.py

묶음(stage) — 오픈 순서
  1 = 서울·경기 시도·시군·구 × 전체 단어 14종
  2 = 동·읍면 × 커튼·블라인드·커튼블라인드
  3 = 동·읍면 × 암막커튼·롤스크린·콤비블라인드
  4 = 동·읍면 × 커튼집·커튼설치·블라인드설치 (새 페이지)
wave = 오픈일(D0)부터 며칠째에 여는지 + 1. 빌드는 OPEN_WAVE 이하만 (scripts/drip/drip_status.py 가 날짜로 계산)
  묶음 1~3 (5,982): 25·50·75·100% 를 D0·D3·D6·D9 에 → wave 1·4·7·10
  묶음 4 (2,088): D1~D30 하루 약 70개 → wave 2~31
표기만 다른 쌍(강남구 커튼 / 강남 커튼, 수원 커튼 / 수원시 …)은 variantOf 로 대표 페이지를 표시한다.
"""
import json, os, re, sys
from collections import Counter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'keywords', '커튼블라인드_전국키워드.xlsx')
SIDOS = ('서울특별시', '경기도')
DONG_W2 = ['커튼', '블라인드', '커튼블라인드']
DONG_W3 = ['암막커튼', '롤스크린', '콤비블라인드']
DONG_W4 = ['커튼집', '커튼설치', '블라인드설치']

allrows = list(load_workbook(SRC, read_only=True)['키워드'].iter_rows(min_row=2, values_only=True))
# 상위 지역은 엑셀 대신 전국 지역 데이터에서 (엑셀은 빈 칸이 잘려 읽히는 경우가 있다)
REG = {r['name']: r for r in json.load(open(os.path.join(ROOT, 'data', 'regions-kr.json'), encoding='utf-8'))['regions']}
# 열: No, 키워드, 단어, 구분, 지역, 지역 단위, 시도, 상위 지역, …
picked, wave_of = [], {}
for r in allrows:
    no, kw, word, grp, region, level, sido, parent = (list(r) + [None] * 8)[:8]
    if sido not in SIDOS:
        continue
    if level in ('시도', '시군', '구'):
        w = 1
    elif word in DONG_W2:
        w = 2
    elif word in DONG_W3:
        w = 3
    elif word in DONG_W4:
        w = 4
    else:
        continue
    wave_of[kw] = w
    parent = (REG.get(region) or {}).get('parent') or parent or region
    picked.append(dict(no=no, kw=kw, slug=kw.replace(' ', '-'), word=word, region=region, level=level, sido=sido, parent=parent, stage=w))
picked.sort(key=lambda p: (p['stage'], p['no']))

# 오픈일 배정 (표기 변형의 대표 페이지는 순서상 항상 앞 → 같은 날이나 먼저 열린다)
old = [p for p in picked if p['stage'] <= 3]
new = [p for p in picked if p['stage'] == 4]
for i, p in enumerate(old):
    p['wave'] = 1 + 3 * (4 * i // len(old))
for i, p in enumerate(new):
    p['wave'] = 2 + 30 * i // len(new)

# 표기 변형 묶음: 같은 시도·단어, 지역명에서 '구/시/군' 과 '서울 ' '경기 ' 접두를 떼면 같은 경우
norm = lambda n: re.sub(r'(구|시|군)$', '', re.sub(r'^(서울|경기) ', '', n))
main = {}
for p in picked:
    key = (p['sido'], p['word'], norm(p['region']), p['level'] in ('동', '읍면'))
    if key in main:
        p['variantOf'] = main[key]
    else:
        main[key] = p['slug']
assert len({p['slug'] for p in picked}) == len(picked)

json.dump({'pages': picked}, open(os.path.join(ROOT, 'data', 'site-pages.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)

# ---- 엑셀 ----
HEAD = PatternFill('solid', fgColor='23483D')
SEL = PatternFill('solid', fgColor='E3ECE8')
wb = Workbook()
c = Counter(p['stage'] for p in picked)
v = sum(1 for p in picked if 'variantOf' in p)
s1 = wb.active; s1.title = '선정 키워드'
s1.append(['순번', '키워드', '페이지 URL', '묶음', '오픈일(D+)', '단어', '지역', '지역 단위', '시도', '상위 지역', '대표 페이지(표기 변형)', '원본 No', '월검색수 PC', '월검색수 모바일'])
for cell in s1[1]: cell.font, cell.fill = Font(bold=True, color='FFFFFF'), HEAD
for i, p in enumerate(picked, 1):
    s1.append([i, p['kw'], f"/{p['slug']}/", p['stage'], p['wave'] - 1, p['word'], p['region'], p['level'], p['sido'], p['parent'], p.get('variantOf', ''), p['no'], None, None])
for col, w in zip('ABCDEFGHIJKLMN', [7, 28, 32, 7, 9, 13, 18, 9, 12, 12, 28, 9, 12, 14]): s1.column_dimensions[col].width = w
s1.freeze_panes = 'A2'; s1.auto_filter.ref = s1.dimensions

s2 = wb.create_sheet('전국 전체(선정 표시)')
s2.append(['No', '키워드', '단어', '구분', '지역', '지역 단위', '시도', '상위 지역', '서울경기 선정', '묶음'])
for cell in s2[1]: cell.font, cell.fill = Font(bold=True, color='FFFFFF'), HEAD
for r in allrows:
    w = wave_of.get(r[1])
    s2.append((list(r) + [None] * 8)[:8] + (['O', w] if w else ['', '']))
    if w:
        for cell in s2[s2.max_row]: cell.fill = SEL
for col, w in zip('ABCDEFGHIJ', [8, 30, 14, 7, 22, 9, 16, 12, 12, 9]): s2.column_dimensions[col].width = w
s2.freeze_panes = 'A2'; s2.auto_filter.ref = s2.dimensions

out = os.path.join(ROOT, 'keywords', '커튼블라인드_서울경기_8000.xlsx')
wb.save(out)
sys.stdout.reconfigure(encoding='utf-8')
print('합계', len(picked), '묶음', dict(sorted(c.items())), '표기 변형', v)
print('오픈일별', dict(sorted(Counter(p['wave'] - 1 for p in picked).items())))
print('->', out)
