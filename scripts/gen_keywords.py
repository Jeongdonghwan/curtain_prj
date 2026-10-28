"""data/regions-kr.json × 커튼·블라인드 단어 → keywords/커튼블라인드_전국키워드.xlsx

  python -I scripts/gen_keywords.py
"""
import json, os, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
regions = json.load(open(os.path.join(ROOT, 'data', 'regions-kr.json'), encoding='utf-8'))['regions']

BASE = ['커튼', '블라인드', '커튼블라인드']                                   # 모든 지역
LOCAL = ['암막커튼', '롤스크린', '콤비블라인드', '커튼집', '커튼설치', '블라인드설치']  # 모든 지역
WIDE = ['우드블라인드', '버티컬블라인드', '전동커튼', '맞춤커튼', '블라인드업체']      # 시도·시군·구만 (검색량 큰 단위)
GROUP = {w: ('기본' if w in BASE else '세부') for w in BASE + LOCAL + WIDE}

rows = []
for r in regions:
    words = BASE + LOCAL + (WIDE if r['level'] in ('시도', '시군', '구') else [])
    for w in words:
        rows.append([f'{r["name"]} {w}', w, GROUP[w], r['name'], r['level'], r['sido'], r['parent']])

kws = [x[0] for x in rows]
assert len(kws) == len(set(kws)), '키워드 중복'

HEAD = PatternFill('solid', fgColor='5B3A29')
wb = Workbook()


def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color='FFFFFF'), HEAD, Alignment(horizontal='center')
    for d in data:
        ws.append(d)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


ws = wb.active
ws.title = '요약'
lv = Counter(r['level'] for r in regions)
kl = Counter(x[4] for x in rows)
summary = [[l, lv[l], kl[l]] for l in ['시도', '시군', '구', '동', '읍면']]
summary += [['합계', len(regions), len(rows)], [], ['단어', '키워드 수', '']]
summary += [[w, n, GROUP[w]] for w, n in Counter(x[1] for x in rows).items()]
sheet(ws, ['지역 단위', '지역 수', '키워드 수'], summary, [16, 12, 12])

sheet(wb.create_sheet('키워드'),
      ['No', '키워드', '단어', '구분', '지역', '지역 단위', '시도', '상위 지역', '월검색수 PC', '월검색수 모바일', '경쟁정도', '비고'],
      [[i + 1] + x + [None, None, None, None] for i, x in enumerate(rows)],
      [8, 30, 14, 7, 22, 9, 16, 12, 12, 14, 10, 16])

sheet(wb.create_sheet('지역목록'), ['No', '지역(검색용)', '지역 단위', '시도', '상위 지역'],
      [[i + 1, r['name'], r['level'], r['sido'], r['parent']] for i, r in enumerate(regions)], [8, 24, 9, 18, 14])

os.makedirs(os.path.join(ROOT, 'keywords'), exist_ok=True)
out = os.path.join(ROOT, 'keywords', '커튼블라인드_전국키워드.xlsx')
wb.save(out)
sys.stdout.reconfigure(encoding='utf-8')
for s in summary[:6]:
    print(*s)
print('keywords', len(rows), '->', out)
